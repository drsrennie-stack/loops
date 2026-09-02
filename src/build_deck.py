#!/usr/bin/env python3
"""Render the retagged slides and assemble the PowerPoint deck.

Every slide gets the same shape: the numbered questions in a column on the
left, the specimen on the right, nothing else. Dr. Rennie's original slides
put the text in whichever corner was free and worded each prompt differently,
so a student had to re-read the layout on every slide. This makes the layout
one thing they never have to think about again.

Type is Plus Jakarta Sans, the display face in the primary palette. Colours
are the primary palette: navy for the prompt, terra for the number and the
eyebrow, white ground.
"""
import os
import sys
from PIL import Image, ImageDraw, ImageFont

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import data  # noqa: E402

ROOT = '/root/t3/out'
PICS = os.path.join(ROOT, 'images/module1-practice')
SLIDES = os.path.join(ROOT, 'slides')
FONT = '/root/t3/fonts/PlusJakartaSans.ttf'

W, H = 1920, 1080
NAVY = (30, 61, 76)
TERRA = (139, 58, 46)
MUTED = (74, 90, 99)
LINE = (215, 221, 224)
WHITE = (255, 255, 255)

# left column
LX, LR = 96, 806          # text runs between these
# picture panel
PX, PY, PR, PB = 862, 74, 1856, 1008


def face(name, size):
    f = ImageFont.truetype(FONT, size)
    f.set_variation_by_name(name)
    return f


def wrap(draw, text, font, width):
    words, lines, cur = text.split(), [], ''
    for w in words:
        trial = (cur + ' ' + w).strip()
        if draw.textlength(trial, font=font) <= width or not cur:
            cur = trial
        else:
            lines.append(cur)
            cur = w
    if cur:
        lines.append(cur)
    return lines


def tracked(draw, xy, text, font, fill, extra):
    """Draw with letter spacing. PIL has no tracking of its own."""
    x, y = xy
    for ch in text:
        draw.text((x, y), ch, font=font, fill=fill)
        x += draw.textlength(ch, font=font) + extra
    return x


def place(canvas, img, box):
    """Fit inside box, centred, with a hairline border."""
    x0, y0, x1, y1 = box
    bw, bh = x1 - x0, y1 - y0
    r = min(bw / img.width, bh / img.height)
    w, h = max(1, int(img.width * r)), max(1, int(img.height * r))
    im = img.resize((w, h), Image.LANCZOS)
    x = x0 + (bw - w) // 2
    y = y0 + (bh - h) // 2
    canvas.paste(im, (x, y))
    ImageDraw.Draw(canvas).rectangle([x - 1, y - 1, x + w, y + h], outline=LINE, width=1)


def title_slide(path):
    c = Image.new('RGB', (W, H), WHITE)
    d = ImageDraw.Draw(c)
    tracked(d, (LX, 300), data.COURSE.upper() + '  ·  ' + data.SCHOOL.upper(),
            face('SemiBold', 26), TERRA, 3)
    y = 360
    for ln in wrap(d, data.TITLE, face('ExtraBold', 108), W - 2 * LX):
        d.text((LX, y), ln, font=face('ExtraBold', 108), fill=NAVY)
        y += 124
    y += 24
    for ln in ['Two questions on every slide. Ninety seconds a slide.',
               'Answer on your sheet by number. The slide moves on without you.']:
        d.text((LX, y), ln, font=face('Regular', 40), fill=MUTED)
        y += 56
    d.text((LX, H - 150), data.BYLINE, font=face('SemiBold', 34), fill=NAVY)
    c.save(path, 'PNG')


def question_slide(stem, qs, path):
    c = Image.new('RGB', (W, H), WHITE)
    d = ImageDraw.Draw(c)

    pic = Image.open(os.path.join(PICS, stem + '.jpg')).convert('RGB')
    place(c, pic, (PX, PY, PR, PB))

    tracked(d, (LX, 96), (data.COURSE + '  ·  Module 1 Practice').upper(),
            face('SemiBold', 22), TERRA, 2.6)

    # Fit both prompts into the column. Shrink together so the two blocks
    # always look like siblings rather than two different sizes.
    top, bottom = 210, H - 130
    for size in range(56, 27, -2):
        nf = face('ExtraBold', size)
        tf = face('SemiBold', size)
        lead = int(size * 1.28)
        blocks, total = [], 0
        for q in qs:
            num = str(q['n']) + '.'
            gap = d.textlength(num + ' ', font=nf)
            lines = wrap(d, q['tag'], tf, (LR - LX) - gap)
            blocks.append((num, gap, lines, nf, tf, lead))
            total += len(lines) * lead
        total += (len(qs) - 1) * int(size * 1.5)
        if total <= bottom - top:
            break

    y = top + max(0, ((bottom - top) - total) // 3)
    for num, gap, lines, nf, tf, lead in blocks:
        d.text((LX, y), num, font=nf, fill=TERRA)
        for i, ln in enumerate(lines):
            d.text((LX + gap, y + i * lead), ln, font=tf, fill=NAVY)
        y += len(lines) * lead + int(lead * 1.18)

    c.save(path, 'PNG')


def main():
    # Wipe first. A previous run numbered the slides differently, and a stale
    # frame left behind gets shipped and read as current.
    if os.path.isdir(SLIDES):
        for f in os.listdir(SLIDES):
            if f.startswith('slide-'):
                os.remove(os.path.join(SLIDES, f))
    os.makedirs(SLIDES, exist_ok=True)
    made = []

    p = os.path.join(SLIDES, 'slide-00-title.png')
    title_slide(p)
    made.append(p)

    for i, (stem, qs) in enumerate(data.slides(), start=1):
        p = os.path.join(SLIDES, 'slide-%02d-q%d-%d.png'
                         % (i, qs[0]['n'], qs[-1]['n']))
        question_slide(stem, qs, p)
        made.append(p)
        print('%-30s q%d and q%d' % (os.path.basename(p), qs[0]['n'], qs[-1]['n']))

    deck(made)
    print('\n%d slides, %d questions' % (len(made), len(data.QUESTIONS)))


def deck(pngs):
    """PowerPoint embeds whatever it is given, so the lossless PNGs go in as
    JPEG. At 1920 wide and quality 92 the type stays clean on a projector and
    the file drops from twenty two megabytes to something that opens quickly
    over a classroom connection."""
    from pptx import Presentation
    from pptx.util import Emu
    import tempfile
    prs = Presentation()
    prs.slide_width = Emu(12192000)          # 13.333in, 16:9
    prs.slide_height = Emu(6858000)          # 7.5in
    blank = prs.slide_layouts[6]
    tmp = tempfile.mkdtemp()
    for p in pngs:
        j = os.path.join(tmp, os.path.basename(p)[:-4] + '.jpg')
        Image.open(p).convert('RGB').save(j, 'JPEG', quality=92, optimize=True)
        s = prs.slides.add_slide(blank)
        s.shapes.add_picture(j, 0, 0, prs.slide_width, prs.slide_height)
    out = os.path.join(ROOT, 'Module-1-Practice-Lab-Exam.pptx')
    prs.save(out)
    print('\nwrote', out, os.path.getsize(out) // 1024, 'KB')


if __name__ == '__main__':
    main()
